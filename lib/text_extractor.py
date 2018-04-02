import docx, pptx, xlrd
import PyPDF2
import re
from os import path
from win32com import client
from zipfile import BadZipFile
from pywintypes import com_error

class TextExtractor():

    def __init__(self):
        self.word = None
        self.excel = None
        self.pwpt = None
        
    def extract_text(self, filepath, password=''):
        filetype = filepath.suffix
    
        if filetype in ['.txt', '.csv', '.xml', '.html', '.htm', '.rtf']:
            return self.extract_plaintext(filepath, password)
        elif filetype == '.docx':
            return self.extract_docx(filepath, password)        
        elif filetype == '.doc':
            return self.open_in_word(filepath, password)
        elif filetype == '.pdf':
            return extract_pdf(filepath, password)
        elif filetype in ['.xls', '.xlsx']:
            return self.extract_xlsx(filepath, password)
        elif filetype == '.pptx':
            return self.extract_pptx(filepath, password)
        elif filetype == '.ppt':
            return open_in_powerpoint(filepath, password)        
        else:
            raise Exception('Document format {} not supported.'.format(filetype))

    def extract_plaintext(self, filepath, password=''):
        try:
            with filepath.open('r', encoding='utf-8-sig') as f:
                return f.read()
        except UnicodeDecodeError:
            with filepath.open('r', encoding='ansi') as f:
                return f.read()

    def extract_pdf(self, filepath, password=''):
        try:
            pdf = filepath.open('rb')
            reader = PyPDF2.PdfFileReader(pdf)
            if reader.isEncrypted and reader.decrypt(password) == 0:
                raise Exception('Incorrect password.')
            full_text = []
            for page in range(reader.numPages):
                p = reader.getPage(page)
                full_text.append(p.extractText())
            text = '\n'.join(full_text)
        except NotImplementedError:
            raise Exception('Acrobat 6.0 encryption not supported.')
        finally:
            pdf.close()
        return text

    def extract_docx(self, filepath, password=''):
        try:
            doc = docx.Document(filepath)
            full_text = []
            for para in doc.paragraphs:
                full_text.append(para.text)
            text = '\n'.join(full_text)
            return text
        except BadZipFile:
            return self.open_in_word(filepath, password)
    
    def open_in_word(self, filepath, password=''):
        if (password == '') or (not isinstance(password, str)):
            password = ' '
        try:
            in_file = path.abspath(filepath)
            if self.word == None:
                self.word = client.DispatchEx('Word.Application')
                self.word.Visible = 0
            doc = self.word.Documents.Open(in_file, 0, 1, 0, password)
            text = doc.Content.Text
            doc.Close()
            text = re.sub('\r', '\n', text)
        except com_error as e:
            if e.hresult == -2147352567:
                raise Exception('Incorrect password.')
            else:
                raise
        return text

    def extract_xlsx(self, filepath, password=''):
        try:
            wb = xlrd.open_workbook(filepath)
            full_text = []
            for sheet in wb.sheets():
                for row in range(sheet.nrows):
                    rowtext = '\t'.join(str(r) for r in sheet.row_values(row) if r != '')
                    full_text.append(rowtext)
            text = '\n'.join(full_text)
        except xlrd.biffh.XLRDError:
            text = self.open_in_excel(filepath, password)
        return text

    def open_in_excel(self, filepath, password=''):
        if (password == '') or (not isinstance(password, str)):
            password = ' '
        try:
            in_file = path.abspath(filepath)
            if self.excel == None:
                self.excel = client.DispatchEx('Excel.Application')
                self.excel.Visible = 0
            wb = self.excel.Workbooks.Open(in_file, 0, 1, None, password)
            full_text = []
            for sheet in wb.Sheets:
                usedrange = sheet.UsedRange()
                if usedrange == None: continue
                for row in usedrange:
                    rowtext = '\t'.join(str(cell) for cell in row if cell != None)
                    full_text.append(rowtext)
            text = '\n'.join(full_text)
            wb.Close()
        except com_error as e:
            if e.hresult == -2147352567:
                raise Exception('Incorrect password.')
            else:
                raise
        return text

    def extract_pptx(self, filepath, password=''):
        try:
            ppt = pptx.Presentation(filepath)
            full_text = []
            for slide in ppt.slides:
                slidetext = '\t'.join(shape.text for shape in slide.shapes if shape.has_text_frame)
                full_text.append(slidetext)
            text = '\n'.join(full_text)
        except Exception:
            text = self.open_in_powerpoint(filepath, password)
        return text

    def open_in_powerpoint(self, filepath, password=''):
        try:
            in_file = path.abspath(filepath)
            if self.pwpt == None:
                self.pwpt = client.DispatchEx('Powerpoint.Application')
            if (password == '') or (not isinstance(password, str)):
                ppt = self.pwpt.Presentations.Open(in_file, 1, 0, 0)
            else:
                window = self.pwpt.ProtectedViewWindows.Open(in_file, password)
                ppt = window.Presentation
            full_text = []
            for slide in ppt.Slides:         
                slidetext = '\t'.join(shape.TextFrame.TextRange.Text for shape in slide.Shapes if shape.HasTextFrame)
                full_text.append(slidetext)
            text = '\n'.join(full_text)
            if (password == '') or (not isinstance(password, str)):
                ppt.Close()
            else:
                window.Close()
        except com_error as e:
            if e.hresult == -2147352567:
                raise Exception('Incorrect password.')
            else:
                raise
        return text

    def cleanup(self):
        if (not self.word == None) and (self.word.Documents.count == 0):
            self.word.Quit()
            self.word = None
        if (not self.excel == None) and (self.excel.Workbooks.count == 0):
            self.excel.Quit()
            self.excel = None
        if (not self.pwpt == None) and (self.pwpt.Presentations.count == 0):
            self.pwpt.Quit()
            self.pwpt = None

if __name__ == '__main__':
    from pathlib import Path
    te = TextExtractor()
    print(te.extract_text(Path('a.pptx'), 'a'))
    te.cleanup()
